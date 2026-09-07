#!/usr/bin/env python3
"""Render Youness Jellouli's resume in the official serif ATS layout."""

from __future__ import annotations

import argparse
import html
import re
import shutil
import subprocess
import sys
from pathlib import Path

try:
    import yaml
except ImportError:
    sys.exit("Missing PyYAML. Run: pip install -r resume/requirements.txt")

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit("Missing python-docx. Run: pip install -r resume/requirements.txt")

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "output"
CONTACT_FILE = ROOT / "profile.yaml"
FONTS = ROOT / "fonts"
INK = "2E3D50"
INK_RGB = RGBColor(0x2E, 0x3D, 0x50)
LINK = "2B7FFF"

LABELS = {
    "fr": {
        "summary": "Résumé",
        "experience": "Expérience",
        "skills": "Compétences",
        "projects": "Projets",
        "education": "Éducation",
        "languages": "Langues",
    },
    "en": {
        "summary": "Summary",
        "experience": "Experience",
        "skills": "Skills",
        "projects": "Projects",
        "education": "Education",
        "languages": "Languages",
    },
}

ICONS = {
    "pin": '<svg viewBox="0 0 24 24" aria-hidden="true"><path d="M12 22s8-6.5 8-12a8 8 0 1 0-16 0c0 5.5 8 12 8 12z"/><circle cx="12" cy="10" r="2.6" fill="#fff"/></svg>',
    "mail": '<svg viewBox="0 0 24 24" aria-hidden="true"><rect x="3" y="5" width="18" height="14" rx="2"/><path d="m4 7 8 6 8-6" fill="none"/></svg>',
    "phone": '<svg viewBox="0 0 24 24" aria-hidden="true"><rect x="8" y="2" width="8" height="20" rx="2"/><circle cx="12" cy="18.2" r="0.9" fill="#fff"/></svg>',
    "linkedin": '<svg viewBox="0 0 24 24" aria-hidden="true"><rect x="2.5" y="2.5" width="19" height="19" rx="2"/><text x="12" y="16.5" text-anchor="middle" font-size="11" font-family="Georgia, serif" font-weight="700" fill="#fff" stroke="none">in</text></svg>',
    "cloud": '<svg viewBox="0 0 24 24" aria-hidden="true"><path d="M7.5 18h9.2A4.3 4.3 0 0 0 18 9.7 6 6 0 0 0 7.2 8.6 4.2 4.2 0 0 0 7.5 18z"/></svg>',
}


def load_yaml(path: Path) -> dict:
    with path.open(encoding="utf-8") as fh:
        data = yaml.safe_load(fh)
    if not isinstance(data, dict):
        raise ValueError(f"Expected a mapping in {path}")
    return data


def lang_of(spec: dict) -> str:
    return (spec.get("meta") or {}).get("language") or "en"


def labels_for(spec: dict) -> dict:
    return LABELS.get(lang_of(spec), LABELS["en"])


def set_run_font(run, name: str = "Times New Roman", size: float = 11, bold: bool = False, italic: bool = False, color: str | None = INK):
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_bottom_border(paragraph, color: str = INK):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def set_paragraph_spacing(paragraph, before=0, after=4, line=1.08):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def add_hyperlink(paragraph, text: str, url: str, size: float = 10, bold: bool = False):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz_cs)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_pr.append(fonts)
    if bold:
        b = OxmlElement("w:b")
        r_pr.append(b)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


TOKEN_RE = re.compile(r"(\*\*.+?\*\*|\[.+?\]\(.+?\))")


def iter_rich(text: str):
    for part in TOKEN_RE.split(str(text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            yield ("bold", part[2:-2])
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            label, url = part[1:-1].split("](", 1)
            yield ("link", label, url)
        else:
            yield ("text", part)


def add_rich_runs(paragraph, text: str, size: float = 10):
    for token in iter_rich(text):
        kind = token[0]
        if kind == "link":
            add_hyperlink(paragraph, token[1], token[2], size=size)
        elif kind == "bold":
            run = paragraph.add_run(token[1])
            set_run_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(token[1])
            set_run_font(run, size=size)


def rich_html(text: str) -> str:
    chunks = []
    for token in iter_rich(text):
        kind = token[0]
        if kind == "link":
            chunks.append(f'<a href="{html.escape(token[2], quote=True)}">{html.escape(token[1])}</a>')
        elif kind == "bold":
            chunks.append(f"<strong>{html.escape(token[1])}</strong>")
        else:
            chunks.append(html.escape(token[1]))
    return "".join(chunks)


def contact_items(contact: dict) -> list[tuple[str, str, str | None]]:
    linkedin = contact.get("linkedin", "")
    linkedin_label = linkedin.replace("https://", "").replace("www.", "")
    if linkedin_label.startswith("linkedin.com/"):
        linkedin_label = linkedin_label.split("linkedin.com/")[-1]
    website = (contact.get("website") or "").replace("https://", "").replace("http://", "")
    phone = contact.get("phone") or ""
    phone_label = phone.replace(" ", "").replace("-", "")
    if phone_label.startswith("+212"):
        rest = phone_label[4:]
        if len(rest) >= 9:
            phone_label = f"(+212){rest[:3]}-{rest[3:]}"
        else:
            phone_label = f"(+212){rest}"
    items = [
        ("pin", contact.get("location") or "", None),
        ("mail", contact.get("email") or "", f"mailto:{contact.get('email')}" if contact.get("email") else None),
        ("phone", phone_label, f"tel:{phone}" if phone else None),
        ("linkedin", linkedin_label, "https://www.linkedin.com/in/youness-jellouli"),
        ("cloud", website, "https://jellyouness.com/"),
    ]
    return [item for item in items if item[1]]


def add_section_heading(doc: Document, text: str):
    p = doc.add_paragraph()
    set_run_font(p.add_run(text.upper()), size=12, bold=True)
    set_paragraph_spacing(p, before=10, after=4, line=1.0)
    add_bottom_border(p)
    return p


def build_docx(spec: dict, contact: dict, dest: Path) -> None:
    labels = labels_for(spec)
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.4)
        section.bottom_margin = Cm(1.2)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        usable = section.page_width - section.left_margin - section.right_margin

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK_RGB
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    heading = doc.styles["Heading 2"]
    heading.font.name = "Times New Roman"
    heading.font.size = Pt(12)
    heading.font.bold = True
    heading.font.color.rgb = INK_RGB
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(4)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(name, before=0, after=4, line=1.0)
    name_run = name.add_run(contact.get("full_name", "Youness JELLOULI"))
    set_run_font(name_run, size=22, bold=True)

    if spec.get("headline") and not (spec.get("meta") or {}).get("hide_headline"):
        h = doc.add_paragraph()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(h, before=0, after=2, line=1.0)
        set_run_font(h.add_run(spec["headline"]), size=11)

    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(contact_p, before=0, after=8, line=1.0)
    bits = []
    for _, label, url in contact_items(contact):
        bits.append((label, url))
    for i, (label, url) in enumerate(bits):
        if i:
            set_run_font(contact_p.add_run("   "), size=10)
        if url:
            add_hyperlink(contact_p, label, url, size=10)
        else:
            set_run_font(contact_p.add_run(label), size=10)

    if spec.get("summary"):
        add_section_heading(doc, labels["summary"])
        s = doc.add_paragraph()
        set_paragraph_spacing(s, before=2, after=2, line=1.15)
        add_rich_runs(s, spec["summary"], size=10)

    if spec.get("experience"):
        add_section_heading(doc, labels["experience"])
        for job in spec["experience"]:
            title = doc.add_paragraph()
            set_paragraph_spacing(title, before=14, after=0, line=1.05)
            set_run_font(title.add_run(job.get("title", "")), size=11, bold=True)

            meta = doc.add_paragraph()
            set_paragraph_spacing(meta, before=0, after=1, line=1.0)
            meta.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
            set_run_font(meta.add_run(job.get("company", "")), size=10, bold=True, italic=True)
            right = ", ".join(part for part in [job.get("dates"), job.get("location")] if part)
            if right:
                set_run_font(meta.add_run("\t" + right), size=10, bold=True, italic=True)

            for bullet in job.get("bullets") or []:
                bp = doc.add_paragraph(style="List Bullet")
                set_paragraph_spacing(bp, before=0, after=0, line=1.08)
                add_rich_runs(bp, str(bullet).strip(), size=10)

    if spec.get("projects"):
        add_section_heading(doc, labels["projects"])
        for project in spec["projects"]:
            header = doc.add_paragraph()
            set_paragraph_spacing(header, before=4, after=0, line=1.05)
            set_run_font(header.add_run(project.get("name", "")), size=11, bold=True)
            if project.get("extra"):
                set_run_font(header.add_run(f"  |  {project['extra']}"), size=10)
            for bullet in project.get("bullets") or []:
                bp = doc.add_paragraph(style="List Bullet")
                set_paragraph_spacing(bp, before=0, after=0, line=1.08)
                add_rich_runs(bp, str(bullet).strip(), size=10)

    if spec.get("skills"):
        add_section_heading(doc, labels["skills"])
        for group in spec["skills"]:
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=5, after=5, line=1.2)
            set_run_font(p.add_run(f"{group['category']}: "), size=10, bold=True)
            set_run_font(p.add_run(", ".join(group.get("items") or [])), size=10)

    if spec.get("education"):
        add_section_heading(doc, labels["education"])
        for edu in spec["education"]:
            p = doc.add_paragraph()
            set_paragraph_spacing(p, before=3, after=0, line=1.05)
            set_run_font(p.add_run(edu.get("degree", "")), size=10, bold=True)
            if edu.get("extra"):
                e = doc.add_paragraph()
                set_paragraph_spacing(e, before=0, after=2, line=1.0)
                set_run_font(e.add_run(edu["extra"]), size=10)

    if spec.get("languages"):
        add_section_heading(doc, labels["languages"])
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=1, after=0)
        set_run_font(p.add_run(" · ".join(spec["languages"])), size=10)

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)


def contact_line(contact: dict) -> str:
    return "  ".join(label for _, label, _ in contact_items(contact))


def build_txt(spec: dict, contact: dict, dest: Path) -> None:
    labels = labels_for(spec)
    plain = lambda t: TOKEN_RE.sub(lambda m: m.group(0).strip("*").split("](")[0].lstrip("["), str(t))
    lines = [contact.get("full_name", "Youness JELLOULI")]
    if spec.get("headline") and not (spec.get("meta") or {}).get("hide_headline"):
        lines.append(spec["headline"])
    lines += [contact_line(contact), ""]
    if spec.get("summary"):
        lines += [labels["summary"].upper(), plain(spec["summary"]), ""]
    if spec.get("experience"):
        lines.append(labels["experience"].upper())
        for job in spec["experience"]:
            lines.append(job.get("title", ""))
            lines.append(f"{job.get('company', '')}    {job.get('dates', '')}, {job.get('location', '')}")
            for bullet in job.get("bullets") or []:
                lines.append(f"• {plain(bullet).strip()}")
            lines.append("")
    if spec.get("projects"):
        lines.append(labels["projects"].upper())
        for project in spec["projects"]:
            extra = f" | {project['extra']}" if project.get("extra") else ""
            lines.append(f"{project.get('name', '')}{extra}")
            for bullet in project.get("bullets") or []:
                lines.append(f"• {plain(bullet).strip()}")
            lines.append("")
    if spec.get("skills"):
        lines.append(labels["skills"].upper())
        for group in spec["skills"]:
            lines.append(f"{group['category']}: {', '.join(group.get('items') or [])}")
        lines.append("")
    if spec.get("education"):
        lines.append(labels["education"].upper())
        for edu in spec["education"]:
            lines.append(edu.get("degree", ""))
            if edu.get("extra"):
                lines.append(edu["extra"])
        lines.append("")
    if spec.get("languages"):
        lines += [labels["languages"].upper(), " · ".join(spec["languages"]), ""]
    dest.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def build_html(spec: dict, contact: dict, dest: Path) -> None:
    labels = labels_for(spec)
    name = html.escape(contact.get("full_name", "Youness JELLOULI"))
    light_font = (FONTS / "Merriweather-latin-300-normal.woff2").as_uri()
    light_italic_font = (FONTS / "Merriweather-latin-300-italic.woff2").as_uri()
    regular_font = (FONTS / "Merriweather-latin-400-normal.woff2").as_uri()
    bold_font = (FONTS / "Merriweather-latin-700-normal.woff2").as_uri()
    italic_font = (FONTS / "Merriweather-latin-400-italic.woff2").as_uri()
    bold_italic_font = (FONTS / "Merriweather-latin-700-italic.woff2").as_uri()

    contact_html = []
    for icon, label, url in contact_items(contact):
        inner = f'{ICONS[icon]}<span>{html.escape(label)}</span>'
        if url:
            contact_html.append(f'<a class="item" href="{html.escape(url, quote=True)}">{inner}</a>')
        else:
            contact_html.append(f'<span class="item">{inner}</span>')

    headline_html = ""
    if spec.get("headline") and not (spec.get("meta") or {}).get("hide_headline"):
        headline_html = f'<p class="headline">{html.escape(spec["headline"])}</p>'

    sections = []
    if spec.get("summary"):
        sections.append(f"<h2>{labels['summary']}</h2><p class='summary'>{rich_html(spec['summary'])}</p>")

    if spec.get("experience"):
        blocks = ["<h2>" + labels["experience"] + "</h2>"]
        for job in spec["experience"]:
            right = ", ".join(part for part in [job.get("dates"), job.get("location")] if part)
            bullets = "".join(f"<li>{rich_html(str(b).strip())}</li>" for b in (job.get("bullets") or []))
            blocks.append(
                f"""<div class="job">
              <p class="job-title">{html.escape(job.get('title', ''))}</p>
              <p class="job-meta"><span class="company">{html.escape(job.get('company', ''))}</span><span class="when">{html.escape(right)}</span></p>
              <ul>{bullets}</ul>
            </div>"""
            )
        sections.append("\n".join(blocks))

    if spec.get("projects"):
        blocks = ["<h2>" + labels["projects"] + "</h2>"]
        for project in spec["projects"]:
            extra = html.escape(project.get("extra") or "")
            extra_html = f'<span class="when">{extra}</span>' if extra else ""
            bullets = "".join(f"<li>{rich_html(str(b).strip())}</li>" for b in (project.get("bullets") or []))
            list_html = f"<ul>{bullets}</ul>" if bullets else ""
            blocks.append(
                f'<div class="job"><p class="job-title">{html.escape(project.get("name", ""))} {extra_html}</p>{list_html}</div>'
            )
        sections.append("\n".join(blocks))

    if spec.get("skills"):
        skill_lines = "".join(
            f"<p class='skill'><strong>{html.escape(g['category'])}:</strong> {html.escape(', '.join(g.get('items') or []))}</p>"
            for g in spec["skills"]
        )
        sections.append(f"<h2>{labels['skills']}</h2>{skill_lines}")

    if spec.get("education"):
        edu_lines = []
        for edu in spec["education"]:
            extra = f"<p class='edu-meta'>{html.escape(edu.get('extra', ''))}</p>" if edu.get("extra") else ""
            edu_lines.append(f"<p class='edu-title'>{html.escape(edu.get('degree', ''))}</p>{extra}")
        sections.append(f"<h2>{labels['education']}</h2>" + "".join(edu_lines))

    if spec.get("languages"):
        sections.append(
            f"<h2>{labels['languages']}</h2><p>{html.escape(' · '.join(spec['languages']))}</p>"
        )

    markup = f"""<!DOCTYPE html>
<html lang="{html.escape(lang_of(spec))}">
<head>
  <meta charset="utf-8" />
  <title>{name}</title>
  <style>
    @font-face {{
      font-family: "Merriweather";
      src: url("{light_font}") format("woff2");
      font-weight: 300;
      font-style: normal;
      font-display: block;
    }}
    @font-face {{
      font-family: "Merriweather";
      src: url("{light_italic_font}") format("woff2");
      font-weight: 300;
      font-style: italic;
      font-display: block;
    }}
    @font-face {{
      font-family: "Merriweather";
      src: url("{regular_font}") format("woff2");
      font-weight: 400;
      font-style: normal;
      font-display: block;
    }}
    @font-face {{
      font-family: "Merriweather";
      src: url("{bold_font}") format("woff2");
      font-weight: 700;
      font-style: normal;
      font-display: block;
    }}
    @font-face {{
      font-family: "Merriweather";
      src: url("{italic_font}") format("woff2");
      font-weight: 400;
      font-style: italic;
      font-display: block;
    }}
    @font-face {{
      font-family: "Merriweather";
      src: url("{bold_italic_font}") format("woff2");
      font-weight: 700;
      font-style: italic;
      font-display: block;
    }}
    @page {{ size: A4; margin: 10.5mm 12.5mm 8mm; }}
    * {{ box-sizing: border-box; }}
    html, body {{
      margin: 0;
      padding: 0;
      background: #fff;
      color: #{INK};
      font-family: "Merriweather", "Liberation Serif", "Times New Roman", Times, serif;
      font-size: 10px;
      font-weight: 300;
      line-height: 1.38;
      word-spacing: 0.12em;
    }}
    h1 {{
      text-align: center;
      font-size: 24px;
      font-weight: 700;
      letter-spacing: 0.15px;
      margin: 0 0 5px;
    }}
    .headline {{
      text-align: center;
      margin: 0 0 4px;
      font-size: 10px;
    }}
    .contact {{
      display: flex;
      justify-content: center;
      align-items: center;
      gap: 14px;
      flex-wrap: nowrap;
      font-size: 9px;
      margin: 0 0 10px;
      white-space: nowrap;
    }}
    .contact .item {{
      display: inline-flex;
      align-items: center;
      gap: 3px;
      color: #{INK};
      text-decoration: none;
      white-space: nowrap;
    }}
    .contact svg {{
      width: 10px;
      height: 10px;
      fill: #{INK};
      stroke: #{INK};
      stroke-width: 1.6;
      flex: 0 0 auto;
    }}
    h2 {{
      font-size: 12px;
      font-weight: 700;
      text-transform: uppercase;
      letter-spacing: 0.3px;
      border-bottom: 1px solid #{INK};
      margin: 10px 0 6px;
      padding: 0 0 1px;
    }}
    p {{ margin: 0 0 5px; }}
    .summary {{ text-align: justify; line-height: 1.42; }}
    a {{ color: #{LINK}; text-decoration: none; }}
    ul {{
      margin: 1px 0 2px;
      padding: 0 0 0 11px;
      list-style: none;
    }}
    li {{
      position: relative;
      margin: 0 0 1px;
      padding-left: 11px;
      text-align: justify;
      line-height: 1.2;
    }}
    li::before {{
      content: "";
      position: absolute;
      left: 1px;
      top: calc((1.2em - 2.25px) / 2);
      width: 2.25px;
      height: 2.25px;
      border-radius: 50%;
      background: #{INK};
    }}
    .job {{ margin: 0 0 12px; }}
    .job-title {{
      font-weight: 700;
      font-size: 11px;
      margin: 0;
    }}
    .job-meta {{
      display: flex;
      justify-content: space-between;
      align-items: baseline;
      gap: 10px;
      margin: 0 0 4px;
      font-size: 10px;
    }}
    .company {{ font-weight: 700; font-style: italic; }}
    .when {{ font-weight: 700; font-style: italic; white-space: nowrap; }}
    .skill {{ margin: 0 0 7px; }}
    .edu-title {{ font-weight: 700; margin: 8px 0 0; }}
    .edu-meta {{ margin: 0; font-weight: 300; }}
    strong {{ font-weight: 700; }}
  </style>
</head>
<body>
  <h1>{name}</h1>
  {headline_html}
  <div class="contact">{''.join(contact_html)}</div>
  {''.join(sections)}
</body>
</html>
"""
    dest.write_text(markup, encoding="utf-8")


def build_pdf(html_path: Path, dest: Path) -> None:
    chrome = shutil.which("google-chrome-stable") or shutil.which("google-chrome") or shutil.which("chromium")
    if not chrome:
        raise RuntimeError("Chrome/Chromium is required to print the PDF")
    dest.parent.mkdir(parents=True, exist_ok=True)
    cmd = [
        chrome,
        "--headless=new",
        "--disable-gpu",
        "--no-sandbox",
        "--disable-dev-shm-usage",
        f"--print-to-pdf={dest}",
        "--no-pdf-header-footer",
        html_path.resolve().as_uri(),
    ]
    subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    try:
        from pypdf import PdfReader

        print(f"PDF pages: {len(PdfReader(dest).pages)}")
    except Exception:
        pass


def default_spec_from_profile(profile: dict) -> dict:
    return load_yaml(ROOT / "jobs" / "baseline.yaml")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate the official-layout resume")
    parser.add_argument("--input", type=Path, help="Tailored YAML spec (resume/jobs/<slug>.yaml)")
    parser.add_argument("--from-profile", action="store_true", help="Build the baseline resume")
    parser.add_argument("--out-name", help="Output filename stem, without extension")
    parser.add_argument("--copy-public-pdf", action="store_true", help="Copy PDF to public/Youness-Jellouli.pdf")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    profile = load_yaml(CONTACT_FILE)
    contact = profile["contact"]

    if args.input:
        spec = load_yaml(args.input)
        slug = args.out_name or spec.get("meta", {}).get("slug") or args.input.stem
    elif args.from_profile:
        spec = default_spec_from_profile(profile)
        slug = args.out_name or "Youness-Jellouli"
    else:
        sys.exit("Provide --input resume/jobs/<slug>.yaml or --from-profile")

    OUTPUT.mkdir(parents=True, exist_ok=True)
    stem = OUTPUT / slug
    docx_path = stem.with_suffix(".docx")
    html_path = stem.with_suffix(".html")
    txt_path = stem.with_suffix(".txt")
    pdf_path = stem.with_suffix(".pdf")

    build_docx(spec, contact, docx_path)
    build_txt(spec, contact, txt_path)
    build_html(spec, contact, html_path)
    build_pdf(html_path, pdf_path)

    if args.copy_public_pdf:
        public_pdf = ROOT.parent / "public" / "Youness-Jellouli.pdf"
        shutil.copyfile(pdf_path, public_pdf)
        print(f"Copied PDF to {public_pdf}")

    print("Wrote:")
    for path in (docx_path, pdf_path, txt_path, html_path):
        print(f"  {path}")


if __name__ == "__main__":
    main()
